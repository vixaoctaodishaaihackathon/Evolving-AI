import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_7_slide_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Premium Color Palette
    BG_DARK = RGBColor(10, 15, 29)
    CARD_BG = RGBColor(19, 27, 46)
    CARD_BORDER = RGBColor(40, 60, 100)
    CYAN_NEON = RGBColor(0, 242, 254)
    BLUE_ELECTRIC = RGBColor(56, 189, 248)
    EMERALD = RGBColor(0, 245, 160)
    AMBER = RGBColor(249, 178, 8)
    CRIMSON = RGBColor(255, 51, 102)
    TEXT_LIGHT = RGBColor(240, 244, 252)
    TEXT_MUTED = RGBColor(148, 163, 184)
    ACCENT_PURPLE = RGBColor(168, 85, 247)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    # =========================================================================
    # SLIDE 1: HERO / TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    badge = add_card(s1, Inches(0.8), Inches(0.75), Inches(4.8), Inches(0.42), bg_color=RGBColor(14, 28, 54), border_color=CYAN_NEON)
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "⚡ ENTERPRISE AI SHOWCASE"
    p_b.font.size = Pt(10.5)
    p_b.font.bold = True
    p_b.font.color.rgb = CYAN_NEON

    tb_title = s1.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(5.8), Inches(2.2))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    p1 = tf_t.paragraphs[0]
    p1.text = "Self-Evolving AI\nSupply Chain Network"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT

    tb_sub = s1.shapes.add_textbox(Inches(0.8), Inches(3.7), Inches(5.8), Inches(1.7))
    tf_s = tb_sub.text_frame
    tf_s.word_wrap = True
    p_s = tf_s.paragraphs[0]
    p_s.text = "Smart Real-Time Logistics Tracking, Early Disruption Alerts & Safe Human Control"
    p_s.font.size = Pt(14)
    p_s.font.color.rgb = TEXT_MUTED

    tb_tech = s1.shapes.add_textbox(Inches(0.8), Inches(5.45), Inches(5.8), Inches(1.3))
    tf_tech = tb_tech.text_frame
    tf_tech.word_wrap = True
    p_tech = tf_tech.paragraphs[0]
    p_tech.text = "CORE PRODUCTION STACK:"
    p_tech.font.size = Pt(11)
    p_tech.font.bold = True
    p_tech.font.color.rgb = BLUE_ELECTRIC
    
    p_tech2 = tf_tech.add_paragraph()
    p_tech2.text = "• LangGraph Multi-Agent Team  • ClickHouse Fast Anomaly Detection\n• Neo4j Deep Supplier Graph    • Apache Kafka Live Streaming\n• NeMo Safety & Spend Rules    • Arize Phoenix Tracing"
    p_tech2.font.size = Pt(11)
    p_tech2.font.color.rgb = TEXT_LIGHT

    hero_img = "d:/Evolving_AI/assets/hero.jpg"
    if os.path.exists(hero_img):
        add_card(s1, Inches(6.8), Inches(0.75), Inches(5.8), Inches(6.0), border_color=CYAN_NEON)
        s1.shapes.add_picture(hero_img, Inches(6.85), Inches(0.80), Inches(5.7), Inches(5.9))

    # =========================================================================
    # SLIDE 2: THE PROBLEM (THE HIDDEN SUPPLIER BLINDSPOT & SURGE COSTS)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)

    tb_h2 = s2.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.85))
    tf_h2 = tb_h2.text_frame
    p_h2 = tf_h2.paragraphs[0]
    p_h2.text = "THE PROBLEM: The Hidden Supplier Blindspot & Cost Spikes"
    p_h2.font.size = Pt(25)
    p_h2.font.bold = True
    p_h2.font.color.rgb = TEXT_LIGHT

    card_p1 = add_card(s2, Inches(0.8), Inches(1.55), Inches(5.6), Inches(2.55))
    tf_p1 = card_p1.text_frame
    tf_p1.word_wrap = True
    pp1 = tf_p1.paragraphs[0]
    pp1.text = "1. The Hidden Supplier Blindspot"
    pp1.font.size = Pt(15)
    pp1.font.bold = True
    pp1.font.color.rgb = CRIMSON
    
    pp1_sub = tf_p1.add_paragraph()
    pp1_sub.text = "• Most companies only know their direct suppliers (Tier-1).\n• They have zero visibility into deeper suppliers (Tier-2 parts & Tier-3 raw materials).\n• When a deep raw-material supplier shuts down, companies find out weeks too late—only after factory assembly lines freeze."
    pp1_sub.font.size = Pt(11.5)
    pp1_sub.font.color.rgb = TEXT_LIGHT

    card_p2 = add_card(s2, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.45))
    tf_p2 = card_p2.text_frame
    tf_p2.word_wrap = True
    pp2 = tf_p2.paragraphs[0]
    pp2.text = "2. Sudden Delays & Runaway Costs"
    pp2.font.size = Pt(15)
    pp2.font.bold = True
    pp2.font.color.rgb = AMBER
    
    pp2_sub = tf_p2.add_paragraph()
    pp2_sub.text = "• Chokepoints (like Suez or Panama Canal) cause sudden +300% shipping delays.\n• Spot freight rates jump +75% within 48 hours during geopolitical disruptions.\n• Late port fees cost companies over $8,500 per day for each delayed shipment."
    pp2_sub.font.size = Pt(11.5)
    pp2_sub.font.color.rgb = TEXT_LIGHT

    graph_img = "d:/Evolving_AI/assets/tier_n_graph.jpg"
    if os.path.exists(graph_img):
        add_card(s2, Inches(6.8), Inches(1.55), Inches(5.7), Inches(5.2), border_color=CRIMSON)
        s2.shapes.add_picture(graph_img, Inches(6.85), Inches(1.60), Inches(5.6), Inches(5.1))

    # =========================================================================
    # SLIDE 3: WHY SIMPLE AI FAILS (SAFETY & FINANCIAL GUARDRAILS)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)

    tb_h3 = s3.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.85))
    tf_h3 = tb_h3.text_frame
    p_h3 = tf_h3.paragraphs[0]
    p_h3.text = "THE SAFETY CHALLENGE: Why Raw AI Cannot Run Supply Chains Alone"
    p_h3.font.size = Pt(25)
    p_h3.font.bold = True
    p_h3.font.color.rgb = TEXT_LIGHT

    card_g1 = add_card(s3, Inches(0.8), Inches(1.55), Inches(5.6), Inches(2.55))
    tf_g1 = card_g1.text_frame
    tf_g1.word_wrap = True
    pg1 = tf_g1.paragraphs[0]
    pg1.text = "1. The AI Mistake & Risk Trap"
    pg1.font.size = Pt(15)
    pg1.font.bold = True
    pg1.font.color.rgb = CRIMSON
    
    pg1_sub = tf_g1.add_paragraph()
    pg1_sub.text = "• Standard AI models make guesses and can hallucinate false information.\n• Giving AI direct purchasing power could result in approving 200% price markups or booking unverified vendors.\n• International trade rules and embargoes require strict, guaranteed checks—not probabilistic guesses."
    pg1_sub.font.size = Pt(11.5)
    pg1_sub.font.color.rgb = TEXT_LIGHT

    card_g2 = add_card(s3, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.45))
    tf_g2 = card_g2.text_frame
    tf_g2.word_wrap = True
    pg2 = tf_g2.paragraphs[0]
    pg2.text = "2. Guaranteed Safety with NeMo Guardrails"
    pg2.font.size = Pt(15)
    pg2.font.bold = True
    pg2.font.color.rgb = EMERALD
    
    pg2_sub = tf_g2.add_paragraph()
    pg2_sub.text = "• Hard Spending Limit: Maximum $50,000 without human approval.\n• Price Spike Protection: Rejects markups higher than 15%.\n• Sanctions Filter: Instantly freezes orders from restricted regions.\n• High-Risk Control: Mandates executive sign-off on critical contracts."
    pg2_sub.font.size = Pt(11.5)
    pg2_sub.font.color.rgb = TEXT_LIGHT

    shield_img = "d:/Evolving_AI/assets/guardrails_shield.jpg"
    if os.path.exists(shield_img):
        add_card(s3, Inches(6.8), Inches(1.55), Inches(5.7), Inches(5.2), border_color=EMERALD)
        s3.shapes.add_picture(shield_img, Inches(6.85), Inches(1.60), Inches(5.6), Inches(5.1))

    # =========================================================================
    # SLIDE 4: THE 5-PILLAR ARCHITECTURE
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)

    tb_h4 = s4.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.85))
    tf_h4 = tb_h4.text_frame
    p_h4 = tf_h4.paragraphs[0]
    p_h4.text = "HOW IT WORKS: 5-Pillar Enterprise Autonomous Engine"
    p_h4.font.size = Pt(25)
    p_h4.font.bold = True
    p_h4.font.color.rgb = TEXT_LIGHT

    pillars = [
        {
            "num": "01",
            "title": "Kafka / Redpanda\nLive Event Stream",
            "points": "• Sub-second data updates\n• GPS container pings\n• Live shipping rate alerts\n• Geopolitical news feed",
            "color": CYAN_NEON
        },
        {
            "num": "02",
            "title": "ClickHouse & Neo4j\nFast & Graph Memory",
            "points": "• ClickHouse: Catches sudden\n  price and delay surges\n• Neo4j: Maps entire deep\n  supplier network links",
            "color": BLUE_ELECTRIC
        },
        {
            "num": "03",
            "title": "LangGraph Network\nSmart AI Team",
            "points": "• Market Watch Agent\n• Supply Chain Graph Agent\n• Deal & Contract Agent\n• Dependable multi-step runs",
            "color": ACCENT_PURPLE
        },
        {
            "num": "04",
            "title": "NeMo Guardrails\nSafety & Human Sign-Off",
            "points": "• $50,000 spend ceiling\n• 15% rate markup limit\n• Blocked region filter\n• Pauses for manager review",
            "color": AMBER
        },
        {
            "num": "05",
            "title": "Meta-Learning\nSelf-Evolution Loop",
            "points": "• Scores every decision\n• Self-tunes strategy rules\n• Stores lessons learned\n• Gets smarter each cycle",
            "color": EMERALD
        }
    ]

    col_width = Inches(2.2)
    gap = Inches(0.18)
    start_left = Inches(0.8)

    for i, p in enumerate(pillars):
        left_pos = start_left + i * (col_width + gap)
        c = add_card(s4, left_pos, Inches(1.65), col_width, Inches(4.3), border_color=p["color"])
        tf_c = c.text_frame
        tf_c.word_wrap = True
        
        pn = tf_c.paragraphs[0]
        pn.text = p["num"]
        pn.font.size = Pt(22)
        pn.font.bold = True
        pn.font.color.rgb = p["color"]
        
        pt = tf_c.add_paragraph()
        pt.text = p["title"]
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_LIGHT
        
        pts = tf_c.add_paragraph()
        pts.text = "\n" + p["points"]
        pts.font.size = Pt(11)
        pts.font.color.rgb = TEXT_MUTED

    banner = add_card(s4, Inches(0.8), Inches(6.15), Inches(11.7), Inches(0.7), bg_color=RGBColor(14, 25, 48), border_color=CYAN_NEON)
    tf_ban = banner.text_frame
    p_ban = tf_ban.paragraphs[0]
    p_ban.text = "🚀 Smart AI + Strict Safety Guardrails: Agents solve complex logistics while strict mathematical rules prevent mistakes."
    p_ban.font.size = Pt(12)
    p_ban.font.bold = True
    p_ban.font.color.rgb = CYAN_NEON

    # =========================================================================
    # SLIDE 5: WORKFLOW IN ACTION & HUMAN-IN-THE-LOOP
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)

    tb_h5 = s5.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.85))
    tf_h5 = tb_h5.text_frame
    p_h5 = tf_h5.paragraphs[0]
    p_h5.text = "WORKFLOW IN ACTION: From Disruption to Human Approval"
    p_h5.font.size = Pt(25)
    p_h5.font.bold = True
    p_h5.font.color.rgb = TEXT_LIGHT

    steps_card = add_card(s5, Inches(0.8), Inches(1.55), Inches(5.6), Inches(5.2))
    tf_st = steps_card.text_frame
    tf_st.word_wrap = True
    
    steps = [
        ("1. Live Threat Detected", "Kafka receives delay ping; ClickHouse catches a +75% shipping rate surge.", CYAN_NEON),
        ("2. Impact Mapped Instantly", "Neo4j traces deep supplier connections; flags $14.7M in factory risk.", CRIMSON),
        ("3. Smart Rerouting Created", "Contract Agent designs an alternate route via Saigon (saving 12 days).", BLUE_ELECTRIC),
        ("4. Safety Gate Triggered", "Order exceeds the $50k spend limit; LangGraph safely pauses for approval.", AMBER),
        ("5. Manager Approves & Executes", "Supervisor clicks 'Authorize'; system immediately places the order.", EMERALD)
    ]

    for j, (title, desc, color) in enumerate(steps):
        p_t = tf_st.paragraphs[0] if j == 0 else tf_st.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = color
        
        p_d = tf_st.add_paragraph()
        p_d.text = f"  {desc}\n"
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED

    hud_img = "d:/Evolving_AI/assets/orchestration_hud.jpg"
    if os.path.exists(hud_img):
        add_card(s5, Inches(6.8), Inches(1.55), Inches(5.7), Inches(5.2), border_color=EMERALD)
        s5.shapes.add_picture(hud_img, Inches(6.85), Inches(1.60), Inches(5.6), Inches(5.1))

    # =========================================================================
    # SLIDE 6: THE SELF-EVOLUTION & LEARNING ENGINE
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)

    tb_h6 = s6.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.85))
    tf_h6 = tb_h6.text_frame
    p_h6 = tf_h6.paragraphs[0]
    p_h6.text = "CONTINUOUS LEARNING: How the AI Network Improves Over Time"
    p_h6.font.size = Pt(25)
    p_h6.font.bold = True
    p_h6.font.color.rgb = TEXT_LIGHT

    card_evo = add_card(s6, Inches(0.8), Inches(1.55), Inches(5.6), Inches(5.2))
    tf_evo = card_evo.text_frame
    tf_evo.word_wrap = True
    
    pe1 = tf_evo.paragraphs[0]
    pe1.text = "1. Clear Performance Score (Reward Function)"
    pe1.font.size = Pt(14)
    pe1.font.bold = True
    pe1.font.color.rgb = CYAN_NEON

    pe1_sub = tf_evo.add_paragraph()
    pe1_sub.text = "Score = (0.6 × Risk Reduced) + (3.5 × Days Saved) - Cost Penalty\nEvaluates every resolved incident against real delivery results."
    pe1_sub.font.size = Pt(11)
    pe1_sub.font.color.rgb = TEXT_LIGHT

    pe2 = tf_evo.add_paragraph()
    pe2.text = "\n2. Adaptive Strategy Memory (evolution_memory.json)"
    pe2.font.size = Pt(14)
    pe2.font.bold = True
    pe2.font.color.rgb = EMERALD

    pe2_sub = tf_evo.add_paragraph()
    pe2_sub.text = "• Supplier Trust Weight: Automatically increases for reliable suppliers.\n• Shipping Speed Weight: Gives higher priority to speed during port blockades.\n• Risk Threshold: Adjusts sensitivity based on supervisor decisions."
    pe2_sub.font.size = Pt(11)
    pe2_sub.font.color.rgb = TEXT_LIGHT

    pe3 = tf_evo.add_paragraph()
    pe3.text = "\n3. Lessons Learned (Episodic Memory)"
    pe3.font.size = Pt(14)
    pe3.font.bold = True
    pe3.font.color.rgb = AMBER

    pe3_sub = tf_evo.add_paragraph()
    pe3_sub.text = "Stores key insights (e.g. 'Always maintain backup suppliers for microcontrollers') and feeds them into future AI planning cycles."
    pe3_sub.font.size = Pt(11)
    pe3_sub.font.color.rgb = TEXT_LIGHT

    evo_img = "d:/Evolving_AI/assets/self_evolution.jpg"
    if os.path.exists(evo_img):
        add_card(s6, Inches(6.8), Inches(1.55), Inches(5.7), Inches(5.2), border_color=EMERALD)
        s6.shapes.add_picture(evo_img, Inches(6.85), Inches(1.60), Inches(5.6), Inches(5.1))

    # =========================================================================
    # SLIDE 7: BUSINESS IMPACT, COMPETITIVE ADVANTAGE & ROADMAP
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)

    tb_h7 = s7.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.85))
    tf_h7 = tb_h7.text_frame
    p_h7 = tf_h7.paragraphs[0]
    p_h7.text = "BUSINESS VALUE: Proven ROI, Key Advantages & Future Roadmap"
    p_h7.font.size = Pt(25)
    p_h7.font.bold = True
    p_h7.font.color.rgb = TEXT_LIGHT

    # Top Metric Cards
    m1 = add_card(s7, Inches(0.8), Inches(1.55), Inches(3.6), Inches(1.6), border_color=EMERALD)
    tf_m1 = m1.text_frame
    p_m1_t = tf_m1.paragraphs[0]
    p_m1_t.text = "MONEY SAVED"
    p_m1_t.font.size = Pt(11)
    p_m1_t.font.color.rgb = TEXT_MUTED
    p_m1_v = tf_m1.add_paragraph()
    p_m1_v.text = "$142,500+"
    p_m1_v.font.size = Pt(28)
    p_m1_v.font.bold = True
    p_m1_v.font.color.rgb = EMERALD
    p_m1_s = tf_m1.add_paragraph()
    p_m1_s.text = "Saved per disruption by avoiding late port fees"
    p_m1_s.font.size = Pt(10)
    p_m1_s.font.color.rgb = TEXT_LIGHT

    m2 = add_card(s7, Inches(4.8), Inches(1.55), Inches(3.6), Inches(1.6), border_color=CYAN_NEON)
    tf_m2 = m2.text_frame
    p_m2_t = tf_m2.paragraphs[0]
    p_m2_t.text = "TIME SAVED"
    p_m2_t.font.size = Pt(11)
    p_m2_t.font.color.rgb = TEXT_MUTED
    p_m2_v = tf_m2.add_paragraph()
    p_m2_v.text = "12 Days Faster"
    p_m2_v.font.size = Pt(28)
    p_m2_v.font.bold = True
    p_m2_v.font.color.rgb = CYAN_NEON
    p_m2_s = tf_m2.add_paragraph()
    p_m2_s.text = "Smart rerouting around congested bottleneck ports"
    p_m2_s.font.size = Pt(10)
    p_m2_s.font.color.rgb = TEXT_LIGHT

    m3 = add_card(s7, Inches(8.8), Inches(1.55), Inches(3.7), Inches(1.6), border_color=AMBER)
    tf_m3 = m3.text_frame
    p_m3_t = tf_m3.paragraphs[0]
    p_m3_t.text = "SAFETY & CONTROL"
    p_m3_t.font.size = Pt(11)
    p_m3_t.font.color.rgb = TEXT_MUTED
    p_m3_v = tf_m3.add_paragraph()
    p_m3_v.text = "100% Compliant"
    p_m3_v.font.size = Pt(28)
    p_m3_v.font.bold = True
    p_m3_v.font.color.rgb = AMBER
    p_m3_s = tf_m3.add_paragraph()
    p_m3_s.text = "Zero unapproved orders above the $50k spend limit"
    p_m3_s.font.size = Pt(10)
    p_m3_s.font.color.rgb = TEXT_LIGHT

    # Bottom Left: Competitive Moat
    comp_card = add_card(s7, Inches(0.8), Inches(3.45), Inches(6.8), Inches(3.4))
    tf_comp = comp_card.text_frame
    tf_comp.word_wrap = True
    p_ct = tf_comp.paragraphs[0]
    p_ct.text = "Why Our System Wins vs Existing Tools"
    p_ct.font.size = Pt(14)
    p_ct.font.bold = True
    p_ct.font.color.rgb = BLUE_ELECTRIC
    
    p_rows = [
        "• Traditional ERP (SAP/Oracle): Slow batch schedules, static data, and reactive phone calls.",
        "• Generic AI Chatbots: Make guesses, risk huge financial mistakes, and have no long-term memory.",
        "• Our Self-Evolving Network: Live data streaming, deep supplier relationship mapping, strict safety rules, and continuous learning."
    ]
    for row in p_rows:
        pr = tf_comp.add_paragraph()
        pr.text = row
        pr.font.size = Pt(11.5)
        pr.font.color.rgb = TEXT_LIGHT

    # Bottom Right: Production Roadmap
    road_card = add_card(s7, Inches(8.0), Inches(3.45), Inches(4.5), Inches(3.4), border_color=ACCENT_PURPLE)
    tf_rd = road_card.text_frame
    tf_rd.word_wrap = True
    p_rt = tf_rd.paragraphs[0]
    p_rt.text = "Enterprise Production Roadmap"
    p_rt.font.size = Pt(14)
    p_rt.font.bold = True
    p_rt.font.color.rgb = ACCENT_PURPLE

    road_items = [
        "1. Air Freight & Multi-Modal Rail Dispatch",
        "2. Global Compliance & Cross-Border Privacy",
        "3. Live SAP & Oracle ERP System Sync",
        "4. Multi-Company AI Negotiation Network"
    ]
    for item in road_items:
        pri = tf_rd.add_paragraph()
        pri.text = f"• {item}"
        pri.font.size = Pt(11.5)
        pri.font.color.rgb = TEXT_LIGHT

    out_path_v2 = "d:/Evolving_AI/Self_Evolving_AI_Supply_Chain_Pitch_v2.pptx"
    prs.save(out_path_v2)
    print(f"Presentation saved successfully to: {out_path_v2}")
    
    out_path = "d:/Evolving_AI/Self_Evolving_AI_Supply_Chain_Pitch.pptx"
    try:
        prs.save(out_path)
        print(f"Also updated primary file: {out_path}")
    except PermissionError:
        print(f"[Note] Primary file '{out_path}' is currently open in PowerPoint. Saved updated presentation as '{out_path_v2}'.")

if __name__ == "__main__":
    create_7_slide_deck()
